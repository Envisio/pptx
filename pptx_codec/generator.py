"""Slides JSON in, .pptx out.

    pptx-codec generate < slides.json > deck.pptx

Input:
    {"font": "Arial", "slides": [{"shapes": [shape, ...]}, ...]}

Every shape has x, y, w, h as fractions of a 20 x 11.25 in (16:9) slide.

    text_box (default)  text, size (pt), color (RRGGBB), bold, caps, align (r | ctr), valign (ctr), line_spacing, inset (pt),
                        fill, line (RRGGBB | tx1 | bg1), line_width (pt),
                        geometry (rounded_rectangle | ellipse | block_arc),
                        adjustment (rounded_rectangle: corner radius 0..0.5;
                                    block_arc: [start_deg, end_deg, thickness 0..0.5], clockwise from 3 o'clock)
    picture             kind: "picture", data (base64 PNG or JPEG), drawn into x, y, w, h as is
    table               kind: "table", columns (fractions of w), rows [[cell, ...], ...], valign (ctr), rule (RRGGBB)
                        borderless native table filling x, y, w, h with equal row heights; a cell takes the text_box text keys
                        (text, size, color, bold, caps, align, line_spacing) or is {} when empty; rule draws a dotted line under every row
    chart               kind: "chart", type (doughnut | bar), data [{label, value, color RRGGBB}, ...]
                        native chart filling x, y, w, h: doughnut has a 62% hole and white seams, clockwise from 12 o'clock;
                        bar is one 100% stacked horizontal bar. No title, legend or axes.
    group               kind: "group", shapes [text_box | picture | group, ...] in slide coordinates; the group's frame is
                        the bounding box of its shapes, so the group itself has no x, y, w, h
"""
import base64
import io
import json
import sys

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE
from pptx.util import Emu, Pt

SLIDE_WIDTH = Emu(18288000)
SLIDE_HEIGHT = Emu(10287000)
DRAWING = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
CHART = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
NAMESPACES = 'xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'

ALIGN = {"l": PP_ALIGN.LEFT, "r": PP_ALIGN.RIGHT, "ctr": PP_ALIGN.CENTER}
VALIGN = {"t": MSO_ANCHOR.TOP, "ctr": MSO_ANCHOR.MIDDLE}
GEOMETRY = {"rounded_rectangle": "roundRect", "ellipse": "ellipse", "block_arc": "blockArc"}
THEME_COLORS = {"tx1": MSO_THEME_COLOR.TEXT_1, "bg1": MSO_THEME_COLOR.BACKGROUND_1}
CHARTS = {"doughnut": XL_CHART_TYPE.DOUGHNUT, "bar": XL_CHART_TYPE.BAR_STACKED_100}
TABLE_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid


def position(shape):
    return (
        Emu(round(shape["x"] * SLIDE_WIDTH)),
        Emu(round(shape["y"] * SLIDE_HEIGHT)),
        Emu(round(shape["w"] * SLIDE_WIDTH)),
        Emu(round(shape["h"] * SLIDE_HEIGHT)),
    )


def paint(color, value):
    if value in THEME_COLORS:
        color.theme_color = THEME_COLORS[value]
    else:
        color.rgb = RGBColor.from_string(value)


def adjustments(shape):
    """Preset geometry guides: block_arc angles are in 60000ths of a degree, everything else in 100000ths."""
    value = shape["adjustment"]
    if shape["geometry"] == "block_arc":
        return {"adj1": value[0] * 60000, "adj2": value[1] * 60000, "adj3": value[2] * 100000}
    return {"adj": value * 100000}


def use_font(presentation, font):
    theme = presentation.slide_master.part.part_related_by(RELATIONSHIP_TYPE.THEME)
    root = etree.fromstring(theme.blob)
    for latin in root.iter(f"{DRAWING}latin"):
        if latin.getparent().tag in (f"{DRAWING}majorFont", f"{DRAWING}minorFont"):
            latin.set("typeface", font)
    theme._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def write_text(frame, shape):
    inset = Pt(shape.get("inset", 2))
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = inset
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = VALIGN[shape.get("valign", "t")]
    frame._bodyPr.set("spcFirstLastPara", "1")
    write_runs(frame, shape)


def write_runs(frame, shape):
    for number, line in enumerate(shape["text"].split("\n")):
        paragraph = frame.paragraphs[0] if number == 0 else frame.add_paragraph()
        paragraph.alignment = ALIGN[shape.get("align", "l")]
        paragraph.line_spacing = shape.get("line_spacing", 1.0)

        run = paragraph.add_run()
        run.text = line
        run.font.size = Pt(shape["size"])
        run.font.bold = shape.get("bold", False)
        if "color" in shape:
            paint(run.font.color, shape["color"])
        run.font._rPr.set("kern", "0")
        if shape.get("caps"):
            run.font._rPr.set("cap", "all")


def draw_text_box(shapes, shape):
    box = shapes.add_textbox(*position(shape))

    if "geometry" in shape:
        geometry = box._element.spPr.prstGeom
        geometry.set("prst", GEOMETRY[shape["geometry"]])
        for name, value in adjustments(shape).items() if "adjustment" in shape else ():
            etree.SubElement(geometry.find(f"{DRAWING}avLst"), f"{DRAWING}gd", name=name, fmla=f"val {round(value)}")

    if "fill" in shape:
        box.fill.solid()
        paint(box.fill.fore_color, shape["fill"])

    if "line" in shape:
        paint(box.line.color, shape["line"])
        box.line.width = Pt(shape.get("line_width", 0.75))
        etree.SubElement(box.line._get_or_add_ln(), f"{DRAWING}round")

    if "text" in shape:
        write_text(box.text_frame, shape)


def draw_picture(shapes, shape):
    image = io.BytesIO(base64.b64decode(shape["data"]))
    shapes.add_picture(image, *position(shape))


def draw_table(shapes, shape):
    x, y, w, h = position(shape)
    table = shapes.add_table(len(shape["rows"]), len(shape["columns"]), x, y, w, h).table
    table.first_row = table.horz_banding = False
    table._tbl.tblPr.find(f"{DRAWING}tableStyleId").text = TABLE_STYLE
    for column, share in zip(table.columns, shape["columns"]):
        column.width = Emu(round(w * share))

    for row, cells in zip(table.rows, shape["rows"]):
        row.height = Emu(round(h / len(shape["rows"])))
        for cell, data in zip(row.cells, cells):
            cell.margin_left = cell.margin_right = cell.margin_top = cell.margin_bottom = Pt(2)
            cell.vertical_anchor = VALIGN[shape.get("valign", "t")]
            if "rule" in shape:
                properties = cell._tc.get_or_add_tcPr()
                for index, side in enumerate(("lnL", "lnR", "lnT", "lnB")):
                    line = f'<a:solidFill><a:srgbClr val="{shape["rule"]}"/></a:solidFill><a:prstDash val="sysDot"/>' if side == "lnB" else "<a:noFill/>"
                    properties.insert(index, etree.fromstring(f'<a:{side} {NAMESPACES} w="12700">{line}</a:{side}>'))
            if "text" in data:
                write_runs(cell.text_frame, data)


def draw_chart(shapes, shape):
    data = CategoryChartData()
    if shape["type"] == "doughnut":
        data.categories = [item["label"] for item in shape["data"]]
        data.add_series("Status", [item["value"] for item in shape["data"]])
    else:
        data.categories = ["Status"]
        for item in shape["data"]:
            data.add_series(item["label"], [item["value"]])

    chart = shapes.add_chart(CHARTS[shape["type"]], *position(shape), data).chart
    chart.has_title = False
    chart.has_legend = False
    plot = chart.plots[0]

    if shape["type"] == "doughnut":
        plot._element.find(f"{CHART}holeSize").set("val", "62")
        for point, item in zip(plot.series[0].points, shape["data"]):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = RGBColor.from_string(item["color"])
            point.format.line.color.rgb = RGBColor(255, 255, 255)
            point.format.line.width = Pt(1.5)
    else:
        plot.gap_width = 0
        plot.overlap = 100
        chart.category_axis.visible = False
        chart.value_axis.visible = False
        chart.value_axis.has_major_gridlines = False
        for series, item in zip(plot.series, shape["data"]):
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = RGBColor.from_string(item["color"])

    # Plot fills the frame, chart area is transparent and borderless.
    plot_area = chart._chartSpace.find(f"{CHART}chart/{CHART}plotArea")
    if plot_area.find(f"{CHART}layout") is not None:
        plot_area.remove(plot_area.find(f"{CHART}layout"))
    plot_area.insert(0, etree.fromstring(
        f'<c:layout {NAMESPACES}><c:manualLayout><c:layoutTarget val="inner"/><c:xMode val="edge"/><c:yMode val="edge"/>'
        '<c:x val="0"/><c:y val="0"/><c:w val="1"/><c:h val="1"/></c:manualLayout></c:layout>'))
    chart._chartSpace.find(f"{CHART}chart").addnext(etree.fromstring(f'<c:spPr {NAMESPACES}><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>'))


def draw(shapes, shape):
    kind = shape.get("kind")
    if kind == "picture":
        draw_picture(shapes, shape)
    elif kind == "chart":
        draw_chart(shapes, shape)
    elif kind == "table":
        draw_table(shapes, shape)
    elif kind == "group":
        group = shapes.add_group_shape()
        for child in shape["shapes"]:
            draw(group.shapes, child)
    else:
        draw_text_box(shapes, shape)


def build(deck, output):
    presentation = Presentation()
    presentation.slide_width = SLIDE_WIDTH
    presentation.slide_height = SLIDE_HEIGHT
    if "font" in deck:
        use_font(presentation, deck["font"])

    blank = presentation.slide_layouts.get_by_name("Blank")
    for slide_data in deck["slides"]:
        slide = presentation.slides.add_slide(blank)
        for shape in slide_data["shapes"]:
            draw(slide.shapes, shape)

    presentation.save(output)

